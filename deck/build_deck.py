"""Builds LingoBuddy.pptx (python-pptx). Run: ../.venv/bin/python build_deck.py"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from PIL import Image

INK, INK2, INK3 = RGBColor(0x0E, 0x10, 0x20), RGBColor(0x4A, 0x4E, 0x66), RGBColor(0x8A, 0x8E, 0xA6)
VIOLET, PURPLE, HEART, OK = RGBColor(0x5B, 0x4B, 0xFF), RGBColor(0x8B, 0x5C, 0xF6), RGBColor(0xFF, 0x4D, 0x6D), RGBColor(0x14, 0xB8, 0x7A)
BG, WHITE, LINE, LAV, SOFT = RGBColor(0xF6, 0xF6, 0xFA), RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0xE7, 0xE7, 0xF0), RGBColor(0xB4, 0xA7, 0xFF), RGBColor(0xC9, 0xCA, 0xE0)
TINT = RGBColor(0xEE, 0xEB, 0xFF)
HEAD, BODY = "Arial", "Calibri"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def bg(slide, color):
    f = slide.background.fill; f.solid(); f.fore_color.rgb = color


def text(slide, x, y, w, h, runs, size=16, color=INK, bold=False, font=BODY, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=None):
    """runs: str | list of paragraphs; a paragraph is str or list of (text, {overrides})."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing: p.space_after = Pt(spacing)
        for t, o in ([(para, {})] if isinstance(para, str) else para):
            r = p.add_run(); r.text = t
            r.font.size = Pt(o.get("size", size)); r.font.bold = o.get("bold", bold)
            r.font.name = o.get("font", font); r.font.color.rgb = o.get("color", color)
            r.font.italic = o.get("italic", False)
    return tb


def box(slide, x, y, w, h, fill=WHITE, line=None, radius=0.12, shadow=False):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.adjustments[0] = min(0.5, radius / min(w, h))
    s.fill.solid(); s.fill.fore_color.rgb = fill
    if line: s.line.color.rgb = line; s.line.width = Pt(1)
    else: s.line.fill.background()
    if not shadow:
        # remove theme shadow
        sp = s._element.spPr; ef = sp.makeelement(qn("a:effectLst"), {}); sp.append(ef)
    s.text_frame.text = ""
    return s


def circle(slide, x, y, d, fill, label, size=18, color=WHITE):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    s.fill.solid(); s.fill.fore_color.rgb = fill; s.line.fill.background()
    sp = s._element.spPr; sp.append(sp.makeelement(qn("a:effectLst"), {}))
    tf = s.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = label; r.font.size = Pt(size); r.font.bold = True; r.font.color.rgb = color; r.font.name = HEAD
    return s


def shot(slide, path, x, y, w=None, h=None, frame=WHITE):
    """Screenshot in a rounded 'window' frame. Give w or h; the other follows the aspect ratio."""
    iw, ih = Image.open(path).size
    if w and not h: h = w * ih / iw
    if h and not w: w = h * iw / ih
    pad = 0.08
    box(slide, x - pad, y - pad, w + 2 * pad, h + 2 * pad, fill=frame, line=LINE if frame == WHITE else None, radius=0.18)
    slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    return w, h


def eyebrow(slide, x, y, label, color=VIOLET):
    text(slide, x, y, 8, 0.3, label.upper(), size=12, color=color, bold=True, font=HEAD)


def title(slide, x, y, w, t, color=INK, size=36):
    text(slide, x, y, w, 1.2, t, size=size, color=color, bold=True, font=HEAD)


def notes(slide, t):
    slide.notes_slide.notes_text_frame.text = t


# 1 — Title
s = prs.slides.add_slide(BLANK); bg(s, INK)
text(s, 0.7, 0.65, 6, 0.4, [[("●  ", {"color": HEART}), ("LingoBuddy", {"bold": True, "color": WHITE})]], size=18, font=HEAD)
text(s, 0.7, 1.7, 6.2, 2.6, [[("New city.", {})], [("New language.", {})], [("No problem.", {"color": LAV})]], size=48, color=WHITE, bold=True, font=HEAD)
text(s, 0.7, 4.35, 5.6, 1, "Your AI-powered local language buddy for navigating everyday life in a new city.", size=18, color=SOFT)
text(s, 0.7, 6.55, 6, 0.4, "TCS Tech Day Hackathon  ·  Live demo: localhost:8000", size=12, color=INK3)
shot(s, "shots/home_c.png", 7.0, 1.5, w=5.7, frame=RGBColor(0x1C, 0x1E, 0x36))
notes(s, "Open with the one-liner: LingoBuddy helps newcomers survive their first week in a new city — not by teaching a language, but by teaching how to live there.")

# 2 — Problem
s = prs.slides.add_slide(BLANK); bg(s, WHITE)
eyebrow(s, 0.7, 0.6, "The problem")
title(s, 0.7, 0.95, 11.5, "Newcomers can't talk to their new city")
text(s, 0.7, 1.85, 11, 0.5, "Students and professionals relocating within India land in a place where the everyday language, phrases and etiquette are new.", size=16, color=INK2)
cards = [
    ("1", "Everyday tasks become stressful", "Hailing an auto, ordering lunch or asking directions turns into guesswork and gestures."),
    ("2", "Textbook language doesn't match the street", "Apps teach “Where is the library?”. Delhi says “Bhaiya, Connaught Place chaloge?”."),
    ("3", "No way to rehearse real situations", "You only practise when it matters — with a driver waiting and a meter running."),
]
for i, (n, h, d) in enumerate(cards):
    x = 0.7 + i * 4.05
    box(s, x, 2.9, 3.75, 3.4, fill=BG, radius=0.2)
    circle(s, x + 0.35, 3.25, 0.7, HEART if i == 0 else VIOLET if i == 1 else PURPLE, n, size=20)
    text(s, x + 0.35, 4.2, 3.1, 0.9, h, size=19, bold=True, font=HEAD)
    text(s, x + 0.35, 5.05, 3.1, 1.2, d, size=14, color=INK2)
notes(s, "Frame the pain around real moments: transport, food, shopping, directions, college.")

# 3 — Insight / positioning
s = prs.slides.add_slide(BLANK); bg(s, BG)
eyebrow(s, 0.7, 0.6, "Our insight")
title(s, 0.7, 0.95, 12, "Newcomers don't need a language course. They need survival skills.")
box(s, 0.7, 2.45, 5.8, 4.35, fill=WHITE, line=LINE, radius=0.2)
text(s, 1.1, 2.8, 5, 0.4, "TRADITIONAL LANGUAGE APPS", size=12, color=INK3, bold=True, font=HEAD)
text(s, 1.1, 3.2, 5, 0.6, "Teach you a language.", size=22, color=INK3, bold=True, font=HEAD)
text(s, 1.1, 4.05, 5, 2.4, ["Generic vocabulary and grammar drills", "Progress measured in XP and streaks", "Same content for every city", "Practice happens in a vacuum"], size=16, color=INK2, spacing=10)
box(s, 6.85, 2.45, 5.8, 4.35, fill=INK, radius=0.2)
text(s, 7.25, 2.8, 5, 0.4, "LINGOBUDDY", size=12, color=LAV, bold=True, font=HEAD)
text(s, 7.25, 3.2, 5.2, 0.6, "Teaches you to live there.", size=22, color=WHITE, bold=True, font=HEAD)
text(s, 7.25, 4.05, 5, 2.4, ["Phrases locals actually use, by situation", "Progress measured in Local Confidence", "City-specific: Delhi's metro, autos, markets", "Rehearse with AI locals before the real thing"], size=16, color=SOFT, spacing=10)
notes(s, "This is the core positioning line — say it word for word.")

# 4 — Hero feature 1: I'm New Here
s = prs.slides.add_slide(BLANK); bg(s, WHITE)
eyebrow(s, 0.7, 0.6, "Hero feature 1  ·  “I'm New Here” mode")
title(s, 0.7, 0.95, 5.4, "A first-week survival pack for your city", size=32)
text(s, 0.7, 2.35, 5, 0.6, "Pick a city and a language. LingoBuddy builds a 7-day plan around the situations you'll actually face.", size=15, color=INK2)
days = [("1", "Greetings"), ("2", "Transport"), ("3", "Food"), ("4", "Shopping"), ("5", "Directions"), ("6", "College & Work"), ("7", "Real Conversation")]
for i, (n, t) in enumerate(days):
    y = 3.3 + i * 0.52
    circle(s, 0.7, y, 0.38, OK if i < 3 else VIOLET if i == 3 else LINE, n, size=12, color=WHITE if i < 4 else INK2)
    text(s, 1.25, y + 0.02, 4, 0.35, [[("Day " + n + "  ", {"color": INK3, "bold": True}), (t, {"bold": True})]], size=15, anchor=MSO_ANCHOR.MIDDLE)
shot(s, "shots/pack_c.png", 6.3, 1.0, w=6.3)
notes(s, "Demo: I'm New Here → Delhi → Hindi / Hinglish → Build My Survival Pack. Delhi is fully built; the architecture supports 10 languages and more cities.")

# 5 — Daily lesson
s = prs.slides.add_slide(BLANK); bg(s, BG)
shot(s, "shots/lesson_c.png", 0.7, 1.0, w=6.3)
eyebrow(s, 7.5, 0.6, "Daily lessons")
title(s, 7.5, 0.95, 5.2, "Real phrases, real situations", size=32)
text(s, 7.5, 2.2, 5.1, 0.5, "Day 2 · Transport: “You need to get to Connaught Place.”", size=15, color=INK2, font=BODY)
box(s, 7.5, 2.85, 5.1, 1.55, fill=WHITE, line=LINE, radius=0.18)
text(s, 7.8, 3.03, 4.6, 0.5, "भैया, कनॉट प्लेस चलोगे?", size=22, bold=True, font="Nirmala UI")
text(s, 7.8, 3.55, 4.6, 0.35, "Bhaiya, Connaught Place chaloge?", size=15, color=VIOLET, bold=True)
text(s, 7.8, 3.93, 4.6, 0.35, "“Could you take me to Connaught Place?”", size=13, color=INK3)
text(s, 7.5, 4.7, 5.1, 2.2, [
    [("Devanagari + pronunciation + ", {}), ("natural meaning", {"bold": True})],
    [("Listen", {"bold": True}), (" at normal or slow speed (browser text-to-speech)", {})],
    [("Practice", {"bold": True}), (" with speech recognition and word-level feedback", {})],
    [("Local tips", {"bold": True}), (", e.g. why you call drivers “bhaiya”", {})],
], size=14, color=INK2, spacing=8)
notes(s, "Click Listen, then Practice — the feedback shows which words were good and which need work.")

# 6 — AI roleplay
s = prs.slides.add_slide(BLANK); bg(s, WHITE)
eyebrow(s, 0.7, 0.6, "AI roleplay")
title(s, 0.7, 0.95, 5.2, "Rehearse with AI locals before the real thing", size=32)
items = [("🛺", "Raju bhaiya", "Auto driver · negotiate a fare"), ("☕", "Priya", "Classmate · chai small talk"), ("🍛", "Sethi ji", "Street food stall · order lunch"), ("🛍️", "Vicky", "Sarojini market · bargain")]
for i, (ic, n, d) in enumerate(items):
    y = 2.45 + i * 0.78
    circle(s, 0.7, y, 0.55, TINT, ic, size=16, color=INK)
    text(s, 1.45, y + 0.02, 4.3, 0.3, n, size=15, bold=True, font=HEAD)
    text(s, 1.45, y + 0.3, 4.3, 0.3, d, size=13, color=INK3)
box(s, 0.7, 5.75, 5.1, 1.0, fill=TINT, radius=0.16)
text(s, 0.95, 5.88, 4.7, 0.8, "Live LLM replies stay in character and coach you after every message. Suggested replies keep beginners moving.", size=13, color=INK2, anchor=MSO_ANCHOR.MIDDLE)
shot(s, "shots/rp_c.png", 6.3, 1.0, w=6.3)
notes(s, "7 scenarios. Replies come from Groq (gpt-oss-120b); if the API is unavailable a scripted engine takes over seamlessly, so the demo never breaks.")

# 7 — Hero feature 2: Local Confidence
s = prs.slides.add_slide(BLANK); bg(s, INK)
eyebrow(s, 0.7, 0.6, "Hero feature 2  ·  Local Confidence", color=LAV)
title(s, 0.7, 0.95, 5.3, "We measure readiness, not XP", color=WHITE, size=32)
text(s, 0.7, 2.3, 5.1, 1.1, "78%", size=72, bold=True, color=WHITE, font=HEAD)
text(s, 0.7, 3.5, 5.1, 0.4, "READY TO EXPLORE DELHI", size=14, bold=True, color=RGBColor(0xFF, 0xB3, 0xC2), font=HEAD)
text(s, 0.7, 4.15, 5.1, 2.6, [
    "Confidence per real-life situation: greetings, transport, food, shopping, directions, college, small talk",
    "Weighted by how often you'll need each one in week one",
    "Rises as you practise phrases and finish conversations",
    "AI insight + “Recommended next” turn weak spots into action",
], size=14, color=SOFT, spacing=9)
shot(s, "shots/confidence_c.png", 6.3, 1.0, w=6.3, frame=RGBColor(0x1C, 0x1E, 0x36))
notes(s, "Demo: open Confidence → read the AI insight → Practice Now → small talk with Priya. Watch the Small Talk bar move.")

# 8 — Everything else
s = prs.slides.add_slide(BLANK); bg(s, WHITE)
eyebrow(s, 0.7, 0.6, "A complete learning loop")
title(s, 0.7, 0.95, 12, "Everything a newcomer needs, in one place")
feats = [
    ("💬", "AI Tutor", "Describe a situation, get the exact words. Remembers the conversation."),
    ("🃏", "Flashcards", "Spaced repetition (SM-2 style) built from your survival pack."),
    ("📖", "Reading", "Beginner Delhi stories with translations and comprehension checks."),
    ("🎧", "Listening", "Hear a phrase, pick its meaning. Slow mode for beginners."),
    ("🎯", "CEFR placement", "5-question level check, A1 → C2 structure."),
    ("🏆", "Progress", "Weekly goal, streak, skills, XP and achievements."),
]
for i, (ic, h, d) in enumerate(feats):
    col, row = i % 3, i // 3
    x, y = 0.7 + col * 4.05, 2.1 + row * 2.45
    box(s, x, y, 3.75, 2.15, fill=BG, radius=0.2)
    circle(s, x + 0.3, y + 0.3, 0.6, WHITE, ic, size=16, color=INK)
    text(s, x + 1.05, y + 0.42, 2.6, 0.4, h, size=18, bold=True, font=HEAD)
    text(s, x + 0.3, y + 1.08, 3.2, 0.95, d, size=13.5, color=INK2)
notes(s, "Mention: voice input and text-to-speech use browser APIs, with text fallback when unavailable.")

# 9 — Architecture
s = prs.slides.add_slide(BLANK); bg(s, BG)
eyebrow(s, 0.7, 0.6, "How it's built")
title(s, 0.7, 0.95, 12, "Simple, fast, and it never breaks in a demo")
layers = [
    ("Frontend", "Single-page app · HTML/CSS/JS, no build step\nReact Bits components ported to vanilla JS\nWeb Speech API: TTS + speech recognition", VIOLET),
    ("Backend", "FastAPI (Python)\nSQLite persistence for learner state\n/api/ai proxy + /api/state", PURPLE),
    ("AI", "Groq · openai/gpt-oss-120b\nAlso supports Claude, OpenAI, Ollama\nOffline engine as automatic fallback", HEART),
]
for i, (h, d, c) in enumerate(layers):
    x = 0.7 + i * 4.15
    box(s, x, 2.2, 3.75, 3.1, fill=WHITE, line=LINE, radius=0.2)
    circle(s, x + 0.3, 2.5, 0.55, c, str(i + 1), size=16)
    text(s, x + 1.0, 2.58, 2.6, 0.4, h, size=20, bold=True, font=HEAD)
    text(s, x + 0.3, 3.35, 3.25, 1.9, d.split("\n"), size=14, color=INK2, spacing=8)
    if i < 2:
        a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + 3.82), Inches(3.55), Inches(0.28), Inches(0.35))
        a.fill.solid(); a.fill.fore_color.rgb = INK3; a.line.fill.background()
box(s, 0.7, 5.7, 11.95, 1.05, fill=INK, radius=0.18)
text(s, 1.05, 5.7, 11.3, 1.05, [[("Reliability by design: ", {"bold": True, "color": WHITE}), ("no signup, a pre-loaded demo learner, curated phrases for known situations, and a scripted AI fallback, so judges never see an error.", {"color": SOFT})]], size=15, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "Adding a city or language = adding content to content.js; the engine is language-agnostic.")

# 10 — Roadmap / close
s = prs.slides.add_slide(BLANK); bg(s, INK)
eyebrow(s, 0.7, 0.6, "What's next", color=LAV)
title(s, 0.7, 0.95, 12, "From Delhi to every city a student moves to", color=WHITE)
road = [("Now", "Delhi · Hindi / Hinglish\n7-day pack, 7 AI roleplays,\nLocal Confidence"),
        ("Next", "Mumbai, Bengaluru, Hyderabad,\nChennai with Marathi, Kannada,\nTelugu and Tamil packs"),
        ("Later", "Server-side pronunciation scoring,\ncampus partnerships, peer\nbuddies who are real locals")]
for i, (h, d) in enumerate(road):
    x = 0.7 + i * 4.05
    box(s, x, 2.2, 3.75, 2.6, fill=RGBColor(0x1C, 0x1E, 0x36), radius=0.2)
    text(s, x + 0.35, 2.5, 3, 0.45, h, size=22, bold=True, color=LAV if i == 0 else WHITE, font=HEAD)
    text(s, x + 0.35, 3.15, 3.1, 1.5, d.split("\n"), size=14, color=SOFT)
text(s, 0.7, 5.2, 12, 1.0, [[("Traditional language apps teach you a language. ", {"color": INK3}), ("LingoBuddy teaches you how to live in a new place.", {"color": WHITE, "bold": True})]], size=22, font=HEAD)
text(s, 0.7, 6.55, 12, 0.4, "Thank you  ·  Live demo at localhost:8000", size=13, color=INK3)
notes(s, "Close on the positioning line, then switch to the live demo.")

prs.save("LingoBuddy.pptx")
print("saved LingoBuddy.pptx", len(prs.slides), "slides")
